import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT = DOCS / "generated"
OUT.mkdir(exist_ok=True)
PROFILE_EXAMPLE = ROOT / "profile_data.example.json"
PROFILE_LOCAL = ROOT / "profile_data.json"


DEFAULT_PROFILE = {
    "name": "Deepan Karthick",
    "email": "deepankarthick212@gmail.com",
    "phone": "9487572163",
    "city_country": "Coimbatore, Tamil Nadu, India",
    "github_url": "https://github.com/darkraigamer212-cmd",
    "linkedin_url": "https://www.linkedin.com/in/deepan-karthick-166735374/",
    "portfolio_url": "https://deepan-karthick-portfolio.industrious-keyboard.workers.dev/",
    "degree": "B.Sc. Computer Science with Artificial Intelligence",
    "college": "Rathinam Global University, Coimbatore",
    "graduation_year": "Expected 2028",
    "experience_dates": "2026-Present",
}


def load_profile():
    if PROFILE_LOCAL.exists():
        with PROFILE_LOCAL.open("r", encoding="utf-8") as fh:
            data = json.load(fh)
        return {**DEFAULT_PROFILE, **data}
    return DEFAULT_PROFILE.copy()


def contact_line(profile):
    return (
        f"{profile['email']} | {profile['phone']} | {profile['city_country']} | "
        f"GitHub: {profile['github_url']} | LinkedIn: {profile['linkedin_url']} | "
        f"Portfolio: {profile['portfolio_url']}"
    )


def render_value(text, profile):
    replacements = {
        "ADD_EMAIL": profile["email"],
        "ADD_PHONE": profile["phone"],
        "ADD_CITY_COUNTRY": profile["city_country"],
        "ADD_GITHUB_URL": profile["github_url"],
        "ADD_LINKEDIN_URL": profile["linkedin_url"],
        "ADD_PORTFOLIO_URL": profile["portfolio_url"],
        "ADD_DEGREE": profile["degree"],
        "ADD_COLLEGE": profile["college"],
        "ADD_GRADUATION_YEAR": profile["graduation_year"],
        "ADD_DATES": profile["experience_dates"],
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text


ATS = {
    "variant": "ats",
    "accent": "1E4A37",
    "filename": "karthik_ats_resume.docx",
    "pdf_filename": "karthik_ats_resume.pdf",
    "subtitle": "B.Sc. CS (AI) Student | React & Python | Business Software | Applied AI",
    "sections": [
        (
            "Summary",
            [
                "Computer Science with Artificial Intelligence student building practical React and JavaScript software for local-business workflows plus defensive, local-first AI, data, cloud, and cybersecurity prototypes. Delivered two working flagship products and a 30-project Applied Labs catalog supported by 142 lab-focused tests; the full repository suite passes 151 tests. Seeking software engineering internships and startup-project opportunities."
            ],
        ),
        (
            "Skills",
            [
                "Languages: Python, JavaScript, C, HTML, CSS; TypeScript basics",
                "Applications and data: React, Vite, responsive UI, dashboards, JSON/CSV validation, data analysis, local-first workflows, PDF/DOCX generation",
                "Engineering: Git, GitHub, automated testing, CI workflows, defensive input validation, REST API basics, Supabase/PostgreSQL basics, cloud/networking fundamentals, technical documentation, AI-assisted development with manual verification",
            ],
        ),
        (
            "Credentials",
            [
                "30 completed credentials across machine learning, generative AI, cloud, networking, programming, data analysis, business, and cybersecurity; certificate archive and working prototypes are linked from the portfolio.",
            ],
        ),
        (
                "Projects",
            [
                "Printing Press ERP | React, TypeScript/JavaScript, Vite, Supabase-compatible architecture",
                "Repaired and deployed a synthetic public demo with owner dashboards, browser-local order/production workflows, persistence, route navigation, and production-priority updates while isolating live client data and authentication.",
                "Verified runtime/store behavior, dashboard routes, SPA deep links, lint, production build, persistence, and a Cloudflare Workers deployment.",
                "Timber CFT Pro with Billing | React, JavaScript, Vite",
                "Converted supplied timber rules into browser-local CFT, ICBM, and raw-millimetre M3 calculations with CFT/M3 pricing, discount, GST, payments, balance/change, strict validation, local history, reopen, and print/Save PDF.",
                "Excluded imported client database files and verified preserved measurement and invoice rules with automated tests.",
                "Applied Labs | React, JavaScript, Python/C artifact generation",
                "Built 30 certificate-connected prototypes across five categories. Selected evidence includes CSV Claim Stress Tester, Detection Contract Drift Guard, Feature Misuse Contract, and Share-Link Afterlife Rehearsal; the full repository suite passes 151 automated tests.",
            ],
        ),
        (
            "Education",
            [
                "ADD_DEGREE, ADD_COLLEGE - ADD_GRADUATION_YEAR | Semester 3 | CGPA 8.0",
                "Relevant coursework: Artificial Intelligence, Programming, Data Structures, Web Development, Database Systems",
            ],
        ),
    ],
}


STARTUP = {
    "variant": "startup",
    "accent": "C4512B",
    "filename": "karthik_startup_resume.docx",
    "pdf_filename": "karthik_startup_resume.pdf",
    "subtitle": "Practical software for real business decisions",
    "sections": [
        (
            "Profile",
            [
                "B.Sc. Computer Science with AI student turning unclear, manual work into small, reviewable software. Current proof: two working business products plus 30 certificate-connected prototypes. I use AI to accelerate planning and debugging while keeping rules, claims, tests, privacy boundaries, and documentation manually reviewable."
            ],
        ),
        (
            "What I Build",
            [
                "React dashboards and internal tools; local-business workflow software; Python/JavaScript automation; data validation; defensive security planning tools; PDF/DOCX output; tested AI-assisted prototypes.",
            ],
        ),
        (
                "Best Proof",
            [
                "01 / Printing Press ERP | React, TypeScript/JavaScript, Vite, Cloudflare",
                "Deployed a synthetic business-workflow demo with owner dashboards, local order/production workflows, persistence, route verification, and an explicit boundary from live client data and authentication.",
                "02 / Timber CFT Pro with Billing | React, JavaScript, deterministic domain modules",
                "Preserved supplied timber measurement rules and added pricing, discount, GST, payments, balance/change, strict validation, local history, reopen, and print/Save PDF without shipping the imported client database.",
                "30 / Applied Labs | Five categories, 142 lab-focused tests",
                "Each certificate maps to a working prototype with a named user, practical problem, inspectable decision logic, reusable output, tests, safety boundaries, and limitations. Selected labs generate claim evidence, regression harnesses, acceptance contracts, and rollback plans.",
            ],
        ),
        (
            "Skills",
            [
                "React, JavaScript, Python, C, HTML, CSS, Vite, Git/GitHub, automated testing, responsive UI, dashboards, business workflows, data validation, defensive security tooling, cloud/networking fundamentals, Supabase/PostgreSQL basics, PDF/DOCX generation, technical documentation, and AI-assisted engineering with manual verification"
            ],
        ),
        (
            "Learning Evidence",
            [
                "Completed 30 credentials across machine learning, generative AI, cloud, networking, programming, data analysis, business, and cybersecurity; each is represented by a working Applied Lab."
            ],
        ),
        (
            "Education",
            ["ADD_DEGREE, ADD_COLLEGE - ADD_GRADUATION_YEAR | Semester 3 | CGPA 8.0"],
        ),
    ],
}


def style_doc(doc, data):
    section = doc.sections[0]
    section.top_margin = Inches(0.58)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(9.35 if data["variant"] == "ats" else 9.45)
    normal.paragraph_format.space_after = Pt(2.5)
    normal.paragraph_format.line_spacing = 1.04

    for style_name, size, color in [
        ("Heading 1", 12.2, data["accent"]),
        ("Heading 2", 11, "111827"),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(2.5)

    bullet = doc.styles["List Bullet"]
    bullet.font.name = "Arial"
    bullet._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    bullet.font.size = Pt(9.2)
    bullet.paragraph_format.left_indent = Inches(0.24)
    bullet.paragraph_format.first_line_indent = Inches(-0.14)
    bullet.paragraph_format.space_after = Pt(1.8)
    bullet.paragraph_format.line_spacing = 1.03


def add_para(doc, text, bold=False, size=10, color="111827", align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(text)
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(9.2)
    run.font.color.rgb = RGBColor.from_string("111827")


def build_resume(data, profile):
    doc = Document()
    style_doc(doc, data)

    name_size = 21 if data["variant"] == "startup" else 19
    add_para(doc, profile["name"], bold=True, size=name_size, color="171713", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, data["subtitle"], size=10.5, color=data["accent"], align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(
        doc,
        contact_line(profile),
        size=8.3,
        color="4B5563",
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    for title, items in data["sections"]:
        doc.add_paragraph(title, style="Heading 1")
        if title in {"Summary", "Profile"}:
            for item in items:
                add_para(doc, render_value(item, profile))
        elif title == "Skills" and len(items) == 1:
            add_para(doc, render_value(items[0], profile))
        else:
            for item in items:
                item = render_value(item, profile)
                if "|" in item:
                    add_para(doc, item, bold=True, color="111827")
                else:
                    add_bullet(doc, item)

    path = OUT / data["filename"]
    doc.save(path)
    return path


def pdf_styles(data):
    base = getSampleStyleSheet()
    accent = colors.HexColor(f"#{data['accent']}")
    return {
        "name": ParagraphStyle(
            "Name",
            parent=base["Title"],
            fontName="Helvetica-Bold",
            fontSize=20,
            leading=24,
            alignment=1,
            textColor=colors.HexColor("#111827"),
            spaceAfter=3,
        ),
        "subtitle": ParagraphStyle(
            "Subtitle",
            parent=base["BodyText"],
            fontName="Helvetica",
            fontSize=10,
            leading=13,
            alignment=1,
            textColor=accent,
            spaceAfter=3,
        ),
        "contact": ParagraphStyle(
            "Contact",
            parent=base["BodyText"],
            fontName="Helvetica",
            fontSize=8,
            leading=10,
            alignment=1,
            textColor=colors.HexColor("#4B5563"),
            spaceAfter=8,
        ),
        "section": ParagraphStyle(
            "Section",
            parent=base["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=12,
            leading=15,
            textColor=accent,
            spaceBefore=8,
            spaceAfter=3,
        ),
        "body": ParagraphStyle(
            "Body",
            parent=base["BodyText"],
            fontName="Helvetica",
            fontSize=9.0,
            leading=10.7,
            textColor=colors.HexColor("#111827"),
            spaceAfter=3,
        ),
        "strong": ParagraphStyle(
            "Strong",
            parent=base["BodyText"],
            fontName="Helvetica-Bold",
            fontSize=9.0,
            leading=10.7,
            textColor=colors.HexColor("#111827"),
            spaceAfter=3,
        ),
        "bullet": ParagraphStyle(
            "Bullet",
            parent=base["BodyText"],
            fontName="Helvetica",
            fontSize=8.8,
            leading=10.4,
            leftIndent=12,
            firstLineIndent=-8,
            textColor=colors.HexColor("#111827"),
            spaceAfter=2,
        ),
    }


def xml_escape(text):
    return (
        str(text)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def build_resume_pdf(data, profile):
    styles = pdf_styles(data)
    path = OUT / data["pdf_filename"]
    doc = SimpleDocTemplate(
        str(path),
        pagesize=LETTER,
        leftMargin=0.65 * inch,
        rightMargin=0.65 * inch,
        topMargin=0.58 * inch,
        bottomMargin=0.58 * inch,
    )
    story = [
        Paragraph(xml_escape(profile["name"]), styles["name"]),
        Paragraph(xml_escape(data["subtitle"]), styles["subtitle"]),
        Paragraph(xml_escape(contact_line(profile)), styles["contact"]),
    ]

    for title, items in data["sections"]:
        story.append(Paragraph(xml_escape(title), styles["section"]))
        for item in items:
            item = render_value(item, profile)
            if title in {"Summary", "Profile"}:
                story.append(Paragraph(xml_escape(item), styles["body"]))
            elif title == "Skills" and len(items) == 1:
                story.append(Paragraph(xml_escape(item), styles["body"]))
            elif "|" in item:
                story.append(Paragraph(xml_escape(item), styles["strong"]))
            else:
                story.append(Paragraph("- " + xml_escape(item), styles["bullet"]))

    doc.build(story)
    return path


if __name__ == "__main__":
    profile = load_profile()
    for payload in (ATS, STARTUP):
        print(build_resume(payload, profile))
        print(build_resume_pdf(payload, profile))
